"""End-to-end API flow against the dev PostGIS database (requires `docker compose up` + migrations + seed)."""

import io
import secrets
from collections.abc import AsyncIterator

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

from app.core.config import get_settings
from app.main import app

pytestmark = pytest.mark.asyncio(loop_scope="session")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="session")
async def client() -> AsyncIterator[httpx.AsyncClient]:
    async with app.router.lifespan_context(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test/api/v1") as c:
            yield c


async def login(client: httpx.AsyncClient, email: str, password: str) -> dict[str, str]:
    res = await client.post("/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


async def _cleanup(tenant_ids: list[str], model_id: str, tag: str) -> None:
    """Remove this run's rows via the owner role (bypasses RLS) so the dev database stays clean."""
    engine = create_async_engine(get_settings().migration_database_url)
    params = {"tenants": tenant_ids, "model": model_id, "supplier": f"%{tag}"}
    device_filter = "SELECT id FROM devices WHERE model_id = CAST(:model AS uuid)"
    async with engine.begin() as conn:
        for stmt in (
            f"DELETE FROM installations WHERE device_id IN ({device_filter})",
            f"DELETE FROM inventory_transactions WHERE device_id IN ({device_filter})",
            "DELETE FROM devices WHERE model_id = CAST(:model AS uuid)",
            "DELETE FROM customers WHERE tenant_id = ANY(CAST(:tenants AS uuid[]))",
            "DELETE FROM users WHERE tenant_id = ANY(CAST(:tenants AS uuid[]))",
            "DELETE FROM tenants WHERE id = ANY(CAST(:tenants AS uuid[]))",
            "DELETE FROM device_models WHERE id = CAST(:model AS uuid)",
            "DELETE FROM suppliers WHERE name LIKE :supplier",
        ):
            await conn.execute(text(stmt), params)
    await engine.dispose()


@pytest.fixture(scope="session")
async def world(client: httpx.AsyncClient) -> AsyncIterator[dict]:
    s = get_settings()
    su = await login(client, s.seed_admin_email, s.seed_admin_password)
    tag = secrets.token_hex(3).upper()
    password = secrets.token_urlsafe(12)

    tenant_a = (await client.post("/tenants", headers=su, json={"name": f"A {tag}", "code": f"A{tag}"})).json()
    tenant_b = (await client.post("/tenants", headers=su, json={"name": f"B {tag}", "code": f"B{tag}"})).json()
    model = (
        await client.post("/device-models", headers=su, json={"brand": "Test", "name": f"M-{tag}", "device_type": "face"})
    ).json()

    users = {}
    for key, tenant, role in (("a", tenant_a, "tenant_admin"), ("b", tenant_b, "tenant_admin"), ("viewer", tenant_a, "viewer")):
        email = f"{key}.{tag.lower()}@example.com"
        res = await client.post(
            "/users",
            headers=su,
            json={"email": email, "full_name": key, "role": role, "password": password, "tenant_id": tenant["id"]},
        )
        assert res.status_code == 201, res.text
        users[key] = await login(client, email, password)

    yield {"su": su, "tag": tag, "tenant_a": tenant_a, "tenant_b": tenant_b, "model": model, **users}
    await _cleanup([tenant_a["id"], tenant_b["id"]], model["id"], tag)


async def test_login_rejects_bad_password(client: httpx.AsyncClient) -> None:
    res = await client.post("/auth/login", json={"email": get_settings().seed_admin_email, "password": "wrong-password"})
    assert res.status_code == 401
    assert res.json()["detail"] == "อีเมลหรือรหัสผ่านไม่ถูกต้อง"


async def test_requires_token(client: httpx.AsyncClient) -> None:
    assert (await client.get("/devices")).status_code == 401


async def test_tenant_cannot_manage_device_models(client: httpx.AsyncClient, world: dict) -> None:
    res = await client.post("/device-models", headers=world["a"], json={"brand": "X", "name": "Y", "device_type": "Z"})
    assert res.status_code == 403


async def test_viewer_is_read_only(client: httpx.AsyncClient, world: dict) -> None:
    res = await client.post(
        "/inventory/check-in", headers=world["viewer"], json={"serial_number": "VIEW-1", "model_id": world["model"]["id"]}
    )
    assert res.status_code == 403


async def test_full_lifecycle_and_tenant_isolation(client: httpx.AsyncClient, world: dict) -> None:
    a, b, su = world["a"], world["b"], world["su"]
    serial = f"T-{world['tag']}-1"

    # Superadmin receives into central stock, then transfers to tenant A.
    res = await client.post("/inventory/check-in", headers=su, json={"serial_number": serial, "model_id": world["model"]["id"]})
    assert res.status_code == 201, res.text
    device = res.json()
    assert device["tenant_id"] is None and device["status"] == "IN_STOCK"

    # Central stock cannot be checked out.
    res = await client.post("/inventory/check-out", headers=su, json={"device_id": device["id"]})
    assert res.status_code == 409

    res = await client.post(
        "/inventory/transfer", headers=su, json={"device_id": device["id"], "target_tenant_id": world["tenant_a"]["id"]}
    )
    assert res.status_code == 200 and res.json()["tenant_name"] == world["tenant_a"]["name"]

    # Tenant B cannot see or move tenant A's device (Row-Level Security).
    assert (await client.get(f"/devices/{device['id']}", headers=b)).status_code == 404
    assert all(d["id"] != device["id"] for d in (await client.get("/devices", headers=b)).json())
    assert (await client.post("/inventory/check-out", headers=b, json={"device_id": device["id"]})).status_code == 404

    # Duplicate serial is rejected even though tenant A cannot see other tenants' rows.
    dup = await client.post("/inventory/check-in", headers=a, json={"serial_number": serial, "model_id": world["model"]["id"]})
    assert dup.status_code == 409

    customer = (await client.post("/customers", headers=a, json={"company_name": f"Cust {world['tag']}"})).json()
    assert customer["tenant_id"] == world["tenant_a"]["id"]

    # Install is only allowed after check-out.
    install_body = {
        "device_id": device["id"],
        "customer_id": customer["id"],
        "install_date": "2026-09-01",
        "latitude": 13.7563,
        "longitude": 100.5018,
        "address": "กรุงเทพฯ",
    }
    assert (await client.post("/installations", headers=a, json=install_body)).status_code == 409
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device["id"]})).status_code == 200
    res = await client.post("/installations", headers=a, json=install_body)
    assert res.status_code == 201, res.text
    installation = res.json()

    # PostGIS nearby: ~1 km away finds it, the other side of the country does not.
    near = (await client.get("/installations/nearby", headers=a, params={"lat": 13.76, "lng": 100.51, "radius_m": 2000})).json()
    assert any(i["id"] == installation["id"] and i["distance_m"] < 2000 for i in near)
    far = (await client.get("/installations/nearby", headers=a, params={"lat": 18.79, "lng": 98.98, "radius_m": 2000})).json()
    assert all(i["id"] != installation["id"] for i in far)
    assert all(i["id"] != installation["id"] for i in (await client.get("/installations", headers=b)).json())

    # Return closes the installation; the device waits for inspection before it is stock again.
    res = await client.post("/inventory/return", headers=a, json={"device_id": device["id"], "note": "เสีย"})
    assert res.status_code == 200 and res.json()["status"] == "UNDER_QC"
    active = (await client.get("/installations", headers=a)).json()
    assert all(i["id"] != installation["id"] for i in active)

    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device["id"]})).json()
    assert [t["transaction_type"] for t in reversed(history)] == ["TRANSFER", "CHECK_OUT", "INSTALL", "RETURN"]
    su_history = (await client.get("/inventory/transactions", headers=su, params={"device_id": device["id"]})).json()
    assert su_history[-1]["transaction_type"] == "CHECK_IN"

    # Dashboard: 30 daily buckets ending today, with today's movements counted; B sees none of A's.
    summary = (await client.get("/dashboard/summary", headers=a)).json()
    activity = summary["activity_30d"]
    assert len(activity) == 30
    assert activity[-1]["counts"].get("INSTALL", 0) >= 1
    assert activity[-1]["total"] == sum(activity[-1]["counts"].values())
    assert summary["recent_transactions"][0]["device_id"] == device["id"]
    b_summary = (await client.get("/dashboard/summary", headers=b)).json()
    assert all(t["device_id"] != device["id"] for t in b_summary["recent_transactions"])


async def test_repair_flow(client: httpx.AsyncClient, world: dict) -> None:
    a, b, su = world["a"], world["b"], world["su"]
    res = await client.post(
        "/inventory/check-in",
        headers=su,
        json={"serial_number": f"R-{world['tag']}", "model_id": world["model"]["id"], "tenant_id": world["tenant_a"]["id"]},
    )
    assert res.status_code == 201, res.text
    device_id = res.json()["id"]
    customer = (await client.post("/customers", headers=a, json={"company_name": f"Repair {world['tag']}"})).json()
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 200
    install = await client.post(
        "/installations",
        headers=a,
        json={"device_id": device_id, "customer_id": customer["id"], "install_date": "2026-09-01", "latitude": 13.7, "longitude": 100.5},
    )
    assert install.status_code == 201, install.text

    # Scanner lookup is case-insensitive and respects tenant isolation.
    found = await client.get(f"/devices/by-serial/r-{world['tag'].lower()}", headers=a)
    assert found.status_code == 200 and found.json()["id"] == device_id
    assert (await client.get(f"/devices/by-serial/R-{world['tag']}", headers=b)).status_code == 404
    assert (await client.get("/devices/by-serial/NO-SUCH-SERIAL", headers=a)).status_code == 404

    # Tenant B cannot touch it; tenant A sends it for repair and the installation closes.
    assert (await client.post("/inventory/send-repair", headers=b, json={"device_id": device_id})).status_code == 404
    res = await client.post("/inventory/send-repair", headers=a, json={"device_id": device_id, "note": "จอไม่ติด"})
    assert res.status_code == 200 and res.json()["status"] == "IN_REPAIR"
    active = (await client.get("/installations", headers=a)).json()
    assert all(i["id"] != install.json()["id"] for i in active)

    # Cannot move while in repair, and QC result is required to finish.
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 409
    assert (await client.post("/inventory/send-repair", headers=a, json={"device_id": device_id})).status_code == 409
    empty_qc = await client.post("/inventory/repair-done", headers=a, json={"device_id": device_id, "qc_note": "   "})
    assert empty_qc.status_code == 422
    res = await client.post("/inventory/repair-done", headers=a, json={"device_id": device_id, "qc_note": "เปลี่ยนจอ ทดสอบผ่าน"})
    assert res.status_code == 200 and res.json()["status"] == "IN_STOCK"

    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device_id})).json()
    assert [t["transaction_type"] for t in history[:2]] == ["REPAIR_DONE", "SEND_REPAIR"]
    assert history[0]["note"] == "เปลี่ยนจอ ทดสอบผ่าน"
    assert history[1]["customer_name"] == customer["company_name"]


async def test_edit_audit_customer_deactivate_and_tx_filters(client: httpx.AsyncClient, world: dict) -> None:
    a, b = world["a"], world["b"]
    res = await client.post("/inventory/check-in", headers=a, json={"serial_number": f"E-{world['tag']}", "model_id": world["model"]["id"]})
    assert res.status_code == 201, res.text
    device_id = res.json()["id"]

    # Editing descriptive fields writes one EDIT row listing what changed; a no-op edit writes nothing.
    res = await client.patch(f"/devices/{device_id}", headers=a, json={"cost": "990.00", "notes": "ชั้น 2"})
    assert res.status_code == 200 and res.json()["status"] == "IN_STOCK"
    await client.patch(f"/devices/{device_id}", headers=a, json={"cost": "990.00"})
    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device_id})).json()
    assert [t["transaction_type"] for t in history] == ["EDIT", "CHECK_IN"]
    assert history[0]["note"] == "แก้ไข: ต้นทุน, หมายเหตุ"
    assert history[0]["from_status"] == history[0]["to_status"] == "IN_STOCK"

    # Filters: type, date range (business days), and tenant isolation still applies.
    edits = (await client.get("/inventory/transactions", headers=a, params={"transaction_type": ["EDIT", "RETIRE"]})).json()
    assert edits and all(t["transaction_type"] in ("EDIT", "RETIRE") for t in edits)
    old = (await client.get("/inventory/transactions", headers=a, params={"date_from": "2020-01-01", "date_to": "2020-01-31"})).json()
    assert old == []
    today = history[0]["occurred_at"][:10]
    ranged = (await client.get("/inventory/transactions", headers=a, params={"date_from": today, "date_to": today})).json()
    assert any(t["device_id"] == device_id for t in ranged)
    reversed_range = {"date_from": "2026-12-31", "date_to": "2026-01-01"}
    assert (await client.get("/inventory/transactions", headers=a, params=reversed_range)).status_code == 422
    assert all(t["device_id"] != device_id for t in (await client.get("/inventory/transactions", headers=b)).json())

    # Customers are deactivated, not deleted, and not while a device is installed there.
    customer = (await client.post("/customers", headers=a, json={"company_name": f"Soft {world['tag']}"})).json()
    assert customer["is_active"] is True
    install_body = {"device_id": device_id, "customer_id": customer["id"], "install_date": "2026-09-01", "latitude": 13.7, "longitude": 100.5}
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 200
    assert (await client.post("/installations", headers=a, json=install_body)).status_code == 201
    res = await client.delete(f"/customers/{customer['id']}", headers=a)
    assert res.status_code == 409
    assert (await client.post("/inventory/return", headers=a, json={"device_id": device_id})).status_code == 200
    assert (await client.post("/inventory/qc-pass", headers=a, json={"device_id": device_id})).status_code == 200
    res = await client.delete(f"/customers/{customer['id']}", headers=a)
    assert res.status_code == 200 and res.json()["is_active"] is False
    listed = {c["id"]: c for c in (await client.get("/customers", headers=a)).json()}
    assert listed[customer["id"]]["is_active"] is False
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 200
    assert (await client.post("/installations", headers=a, json=install_body)).status_code == 422
    res = await client.patch(f"/customers/{customer['id']}", headers=a, json={"is_active": True})
    assert res.status_code == 200 and res.json()["is_active"] is True


async def test_loan_and_qc_flow(client: httpx.AsyncClient, world: dict) -> None:
    a, b = world["a"], world["b"]
    res = await client.post("/inventory/check-in", headers=a, json={"serial_number": f"L-{world['tag']}", "model_id": world["model"]["id"]})
    device_id = res.json()["id"]

    # Due date is required and cannot be in the past; tenant B cannot lend A's device.
    assert (await client.post("/inventory/loan", headers=a, json={"device_id": device_id})).status_code == 422
    past = {"device_id": device_id, "due_date": "2020-01-01"}
    assert (await client.post("/inventory/loan", headers=a, json=past)).status_code == 422
    body = {"device_id": device_id, "due_date": "2099-12-31", "note": "ทดลองใช้"}
    assert (await client.post("/inventory/loan", headers=b, json=body)).status_code == 404
    res = await client.post("/inventory/loan", headers=a, json=body)
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "ON_LOAN" and res.json()["loan_due_date"] == "2099-12-31"
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 409

    # Return → UNDER_QC (due date cleared) → fail needs a reason → IN_REPAIR → repaired → return path again → pass.
    res = await client.post("/inventory/return", headers=a, json={"device_id": device_id})
    assert res.json()["status"] == "UNDER_QC" and res.json()["loan_due_date"] is None
    assert (await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})).status_code == 409
    assert (await client.post("/inventory/qc-fail", headers=a, json={"device_id": device_id, "note": " "})).status_code == 422
    res = await client.post("/inventory/qc-fail", headers=a, json={"device_id": device_id, "note": "ปุ่มกดไม่ติด"})
    assert res.json()["status"] == "IN_REPAIR"
    await client.post("/inventory/repair-done", headers=a, json={"device_id": device_id, "qc_note": "เปลี่ยนปุ่ม"})
    await client.post("/inventory/check-out", headers=a, json={"device_id": device_id})
    await client.post("/inventory/return", headers=a, json={"device_id": device_id})
    res = await client.post("/inventory/qc-pass", headers=a, json={"device_id": device_id, "note": "ปกติ"})
    assert res.status_code == 200 and res.json()["status"] == "IN_STOCK"

    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device_id})).json()
    assert [t["transaction_type"] for t in reversed(history)] == [
        "CHECK_IN", "LOAN", "RETURN", "QC_FAIL", "REPAIR_DONE", "CHECK_OUT", "RETURN", "QC_PASS",
    ]
    assert history[-2]["note"].startswith("ครบกำหนดคืน 31/12/2099")


async def _owner_exec(sql: str, params: dict) -> None:
    """Backdate rows the API never lets you write (past due dates, old movements)."""
    engine = create_async_engine(get_settings().migration_database_url)
    async with engine.begin() as conn:
        await conn.execute(text(sql), params)
    await engine.dispose()


async def test_dashboard_backlog(client: httpx.AsyncClient, world: dict) -> None:
    a, b, tag = world["a"], world["b"], world["tag"]
    ids = {}
    for key in ("qc", "loan", "repair"):
        res = await client.post("/inventory/check-in", headers=a, json={"serial_number": f"BL-{tag}-{key}", "model_id": world["model"]["id"]})
        ids[key] = res.json()["id"]
    await client.post("/inventory/check-out", headers=a, json={"device_id": ids["qc"]})
    await client.post("/inventory/return", headers=a, json={"device_id": ids["qc"]})
    await client.post("/inventory/loan", headers=a, json={"device_id": ids["loan"], "due_date": "2099-01-01"})
    await client.post("/inventory/send-repair", headers=a, json={"device_id": ids["repair"]})

    await _owner_exec(
        "UPDATE devices SET loan_due_date = current_date - 3 WHERE id = CAST(:id AS uuid)", {"id": ids["loan"]}
    )
    await _owner_exec(
        "UPDATE inventory_transactions SET occurred_at = occurred_at - interval '20 days' WHERE device_id = CAST(:id AS uuid)",
        {"id": ids["repair"]},
    )
    # An EDIT today must not restart the repair clock.
    await client.patch(f"/devices/{ids['repair']}", headers=a, json={"notes": "รออะไหล่"})

    summary = (await client.get("/dashboard/summary", headers=a)).json()
    assert summary["pending_qc"] >= 1
    assert ids["loan"] in [d["id"] for d in summary["loan_overdue"]]
    aging = {x["device"]["id"]: x["days"] for x in summary["repair_aging"]}
    assert aging.get(ids["repair"], 0) >= 20

    other = (await client.get("/dashboard/summary", headers=b)).json()
    assert ids["loan"] not in [d["id"] for d in other["loan_overdue"]]
    assert ids["repair"] not in [x["device"]["id"] for x in other["repair_aging"]]


async def test_reports(client: httpx.AsyncClient, world: dict) -> None:
    a, b, tag, model_id = world["a"], world["b"], world["tag"], world["model"]["id"]
    for i, cost in enumerate((1000, 2500)):
        await client.post(
            "/inventory/check-in", headers=a, json={"serial_number": f"RP-{tag}-{i}", "model_id": model_id, "cost": cost}
        )

    def rows_for(report: list[dict], tenant_id: str) -> list[dict]:
        return [r for r in report if r["model_id"] == model_id and r["tenant_id"] == tenant_id]

    balance = (await client.get("/reports/stock-balance", headers=a)).json()
    stocked = [r for r in rows_for(balance, world["tenant_a"]["id"]) if r["status"] == "IN_STOCK"]
    assert len(stocked) == 1 and stocked[0]["count"] >= 2 and float(stocked[0]["total_cost"]) >= 3500
    assert all(r["status"] != "RETIRED" for r in balance)
    assert not rows_for((await client.get("/reports/stock-balance", headers=b)).json(), world["tenant_a"]["id"])

    aging = (await client.get("/reports/aging", headers=a, params={"status": "IN_STOCK"})).json()
    fresh = [x for x in aging if x["device"]["serial_number"] == f"RP-{tag}-0"]
    assert fresh and fresh[0]["days"] == 0
    assert all(x["device"]["status"] == "IN_STOCK" for x in aging)
    assert all(x["device"]["tenant_id"] != world["tenant_a"]["id"] for x in (await client.get("/reports/aging", headers=b)).json())


async def test_device_paging(client: httpx.AsyncClient, world: dict) -> None:
    a, b, tag, model_id = world["a"], world["b"], world["tag"], world["model"]["id"]
    for i in range(3):
        await client.post("/inventory/check-in", headers=a, json={"serial_number": f"PG-{tag}-{i}", "model_id": model_id})

    search = {"search": f"PG-{tag}"}
    res = await client.get("/devices", headers=a, params={**search, "sort": "-serial_number", "skip": 1, "take": 1})
    assert res.status_code == 200 and res.headers["X-Total-Count"] == "3"
    assert [d["serial_number"] for d in res.json()] == [f"PG-{tag}-1"]

    everything = await client.get("/devices", headers=a, params=search)
    assert len(everything.json()) == 3 and everything.headers["X-Total-Count"] == "3"
    assert (await client.get("/devices", headers=b, params=search)).headers["X-Total-Count"] == "0"
    assert (await client.get("/devices", headers=a, params={"sort": "tenant_id"})).status_code == 422

    counts = {c["key"]: c["count"] for c in (await client.get("/devices/status-counts", headers=a)).json()}
    assert counts["IN_STOCK"] >= 3


def _xlsx(rows: list[list[object]]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def test_import_devices(client: httpx.AsyncClient, world: dict) -> None:
    a, tag = world["a"], world["tag"]
    model = f"Test {world['model']['name']}"
    header = ["ซีเรียล", "รุ่น", "MAC", "วันที่ซื้อ", "ต้นทุน", "หมดประกัน", "หมายเหตุ"]
    bad = _xlsx(
        [
            header,
            [f"imp-{tag}-1", model, "00:11:22:33:44:55", "29/09/2569", "12,500", "2027-09-29", "ok"],
            [f"IMP-{tag}-1", model, None, None, None, None, None],
            [f"IMP-{tag}-2", "No Such Model", "zz", None, -5, None, None],
            [None, None, None, None, None, None, None],
        ]
    )
    files = {"file": ("devices.xlsx", bad, XLSX)}

    assert (await client.post("/inventory/import", headers=world["viewer"], files=files)).status_code == 403

    res = await client.post("/inventory/import", headers=a, files=files, params={"dry_run": "true"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["total"], body["valid"], body["invalid"], body["committed"]) == (3, 1, 2, False)
    first, dup, broken = body["rows"]
    assert first["serial_number"] == f"IMP-{tag}-1" and first["purchase_date"] == "2026-09-29" and first["cost"] == "12500.00"
    assert dup["row"] == 3 and "ซ้ำกับแถว 2" in dup["errors"][0]
    assert len(broken["errors"]) == 3

    # Commit refuses while any row is invalid, and nothing is written.
    res = await client.post("/inventory/import", headers=a, files=files, params={"dry_run": "false"})
    assert res.status_code == 422
    assert (await client.get(f"/devices/by-serial/IMP-{tag}-1", headers=a)).status_code == 404

    csv_body = "\n".join([",".join(header), f"IMP-{tag}-1,{model},,,,,", f"IMP-{tag}-2,{world['model']['name']},,,,,"])
    files = {"file": ("devices.csv", csv_body.encode("utf-8-sig"), "text/csv")}
    res = await client.post("/inventory/import", headers=a, files=files, params={"dry_run": "false"})
    assert res.status_code == 200, res.text
    assert res.json()["committed"] is True and res.json()["valid"] == 2
    device = (await client.get(f"/devices/by-serial/IMP-{tag}-2", headers=a)).json()
    assert device["tenant_id"] == world["tenant_a"]["id"] and device["status"] == "IN_STOCK"
    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device["id"]})).json()
    assert [t["transaction_type"] for t in history] == ["CHECK_IN"]

    # Re-importing the same file now fails validation on existing serials.
    res = await client.post("/inventory/import", headers=a, files=files)
    assert res.json()["invalid"] == 2

    template = await client.get("/inventory/import/template", headers=a)
    assert template.status_code == 200 and template.headers["content-type"] == XLSX
    sheet = load_workbook(io.BytesIO(template.content)).worksheets[0]
    assert [c.value for c in sheet[1]] == header


async def test_change_own_password(client: httpx.AsyncClient, world: dict) -> None:
    email = f"pw.{world['tag'].lower()}@example.com"
    old, new = "old-password-123", "new-password-456"
    res = await client.post(
        "/users",
        headers=world["su"],
        json={"email": email, "full_name": "pw", "role": "staff", "password": old, "tenant_id": world["tenant_a"]["id"]},
    )
    assert res.status_code == 201, res.text
    headers = await login(client, email, old)

    res = await client.post("/auth/change-password", headers=headers, json={"current_password": "wrong-pass", "new_password": new})
    assert res.status_code == 422
    assert res.json()["detail"] == "รหัสผ่านปัจจุบันไม่ถูกต้อง"
    res = await client.post("/auth/change-password", headers=headers, json={"current_password": old, "new_password": old})
    assert res.status_code == 422
    res = await client.post("/auth/change-password", headers=headers, json={"current_password": old, "new_password": "short"})
    assert res.status_code == 422

    res = await client.post("/auth/change-password", headers=headers, json={"current_password": old, "new_password": new})
    assert res.status_code == 204, res.text
    await login(client, email, new)
    assert (await client.post("/auth/login", json={"email": email, "password": old})).status_code == 401


async def test_tenant_admin_cannot_create_users_in_other_tenant(client: httpx.AsyncClient, world: dict) -> None:
    res = await client.post(
        "/users",
        headers=world["a"],
        json={
            "email": f"x.{world['tag'].lower()}@example.com",
            "full_name": "x",
            "role": "superadmin",
            "password": "long-enough-password",
        },
    )
    assert res.status_code == 403


async def test_suppliers_crud_and_send_repair(client: httpx.AsyncClient, world: dict) -> None:
    a, su, tag = world["a"], world["su"], world["tag"]
    body = {"name": f"Repair Shop {tag}", "phone": "02-000-0000", "email": f"shop.{tag.lower()}@example.com"}

    assert (await client.post("/suppliers", headers=a, json=body)).status_code == 403
    res = await client.post("/suppliers", headers=su, json=body)
    assert res.status_code == 201, res.text
    supplier = res.json()
    dup = await client.post("/suppliers", headers=su, json={"name": f"repair shop {tag}"})
    assert dup.status_code == 409 and dup.json()["detail"] == "ชื่อผู้จำหน่าย/ผู้ซ่อมนี้มีอยู่แล้ว"

    res = await client.patch(f"/suppliers/{supplier['id']}", headers=su, json={"contact_person": "คุณช่าง"})
    assert res.status_code == 200 and res.json()["contact_person"] == "คุณช่าง"
    # Tenants read the shared list so they can pick a repair shop.
    assert any(x["id"] == supplier["id"] for x in (await client.get("/suppliers", headers=a)).json())

    device = (
        await client.post("/inventory/check-in", headers=a, json={"serial_number": f"SUP-{tag}", "model_id": world["model"]["id"]})
    ).json()
    missing = await client.post(
        "/inventory/send-repair",
        headers=a,
        json={"device_id": device["id"], "supplier_id": "00000000-0000-0000-0000-000000000000"},
    )
    assert missing.status_code == 404
    res = await client.post(
        "/inventory/send-repair", headers=a, json={"device_id": device["id"], "supplier_id": supplier["id"], "note": "จอเสีย"}
    )
    assert res.status_code == 200 and res.json()["status"] == "IN_REPAIR"
    history = (await client.get("/inventory/transactions", headers=a, params={"device_id": device["id"]})).json()
    assert history[0]["transaction_type"] == "SEND_REPAIR" and history[0]["supplier_name"] == body["name"]

    res = await client.delete(f"/suppliers/{supplier['id']}", headers=su)
    assert res.status_code == 409
    spare = (await client.post("/suppliers", headers=su, json={"name": f"Spare {tag}"})).json()
    assert (await client.delete(f"/suppliers/{spare['id']}", headers=su)).status_code == 204
