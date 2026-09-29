"""Read uploaded .xlsx/.csv files into header→value rows, and build the import template."""

import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.errors import ValidationError

Row = tuple[int, dict[str, Any]]


def _decode_csv(content: bytes) -> str:
    # Excel on Thai Windows saves CSV as cp874 unless "UTF-8" is chosen explicitly.
    for encoding in ("utf-8-sig", "cp874"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationError("อ่านไฟล์ CSV ไม่ได้ กรุณาบันทึกเป็น UTF-8")


def _rows(values: list[list[Any]]) -> list[Row]:
    if not values:
        return []
    headers = [str(h).strip() if h is not None else "" for h in values[0]]
    result: list[Row] = []
    for index, raw in enumerate(values[1:], start=2):
        cells = {h: v for h, v in zip(headers, raw, strict=False) if h}
        if any(v not in (None, "") and str(v).strip() for v in cells.values()):
            result.append((index, cells))
    return result


def read_table(content: bytes, filename: str) -> list[Row]:
    """First sheet (or CSV) with the header on row 1; blank rows are skipped. Row numbers match the file."""
    name = filename.lower()
    if name.endswith(".csv"):
        return _rows([list(r) for r in csv.reader(io.StringIO(_decode_csv(content)))])
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError("อ่านไฟล์ Excel ไม่ได้ ไฟล์อาจเสียหายหรือไม่ใช่ .xlsx") from exc
        sheet = book.worksheets[0]
        rows = _rows([list(r) for r in sheet.iter_rows(values_only=True)])
        book.close()
        return rows
    raise ValidationError("รองรับเฉพาะไฟล์ .xlsx หรือ .csv")


def build_template(headers: list[str], example: list[Any], models: list[str]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "อุปกรณ์"
    sheet.append(headers)
    sheet.append(example)
    header_fill = PatternFill("solid", fgColor="DDE3FF")
    for col, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        sheet.column_dimensions[get_column_letter(col)].width = 22
    sheet.freeze_panes = "A2"

    ref = book.create_sheet("รุ่นที่ใช้ได้")
    ref.append(["ชื่อรุ่น (คัดลอกไปใส่คอลัมน์ รุ่น)"])
    ref.cell(row=1, column=1).font = Font(bold=True)
    ref.column_dimensions["A"].width = 40
    for model in models:
        ref.append([model])

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
